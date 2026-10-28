import csv
import os
from datetime import datetime

from colorama import Fore, init
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from procesamiento import calcular_metricas_produccion, limpiar_y_filtrar_lecturas

init(autoreset=True) 

CARPETA = "Datos_Planta"

def main():
    log_crudo = [
        "120.5", "130", "error_sensor", "125.3", "nan",
        "inf", "118.9", "-inf", "135.0", "sin_senal", "128.4",
    ]

    os.makedirs(CARPETA, exist_ok=True)

    validos, incidencias = limpiar_y_filtrar_lecturas(log_crudo)
    print(Fore.GREEN + f"Lecturas validas: {len(validos)}")
    print(Fore.YELLOW + f"Incidencias detectadas: {len(incidencias)}")

    ruta_bitacora = os.path.join(CARPETA, "bitacora_alertas.txt")
    with open(ruta_bitacora, "a", encoding="utf-8") as archivo:
        marca = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        for incidencia in incidencias:
            archivo.write(f"[{marca}] {incidencia}\n")
            print(Fore.RED + f"ALERTA: {incidencia}")

    minimo, maximo, promedio = calcular_metricas_produccion(validos)
    ruta_csv = os.path.join(CARPETA, "reporte_metricas.csv")
    with open(ruta_csv, "w", newline="", encoding="utf-8") as archivo:
        escritor = csv.writer(archivo, delimiter=";")
        escritor.writerow(["Metrica", "Valor"])
        escritor.writerow(["Minimo", minimo])
        escritor.writerow(["Maximo", maximo])
        escritor.writerow(["Promedio", round(promedio, 2) if promedio is not None else ""])
    print(Fore.CYAN + f"CSV generado: {ruta_csv}")

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Produccion_Mensual"

    hoja["A1"] = "Codigo Componente"
    hoja["B1"] = "Cantidad Producida"
    relleno = PatternFill(start_color="000080", end_color="000080", fill_type="solid")
    fuente = Font(bold=True, color="FFFFFF")
    for celda in (hoja["A1"], hoja["B1"]):
        celda.font = fuente
        celda.fill = relleno

    datos = [
        ("COMP-001", 1500),
        ("COMP-002", 1320),
        ("COMP-003", 1745),
        ("COMP-004", 980),
        ("COMP-005", 2010),
        ("COMP-006", 1630),
    ]
    for fila in datos:
        hoja.append(fila)

    hoja.column_dimensions["A"].width = 22
    hoja.column_dimensions["B"].width = 22

    ruta_xlsx = os.path.join(CARPETA, "Reporte_Calidad.xlsx")
    libro.save(ruta_xlsx)
    print(Fore.CYAN + f"Excel generado: {ruta_xlsx}")

    libro_aud = load_workbook(ruta_xlsx)
    hoja_aud = libro_aud["Produccion_Mensual"]
    valor_original = hoja_aud["B2"].value
    hoja_aud["B2"] = 1480  
    libro_aud.save(ruta_xlsx)
    print(
        Fore.MAGENTA
        + f"Auditoria: B2 corregida de {valor_original} a {hoja_aud['B2'].value}"
    )


if __name__ == "_main_":
    main()